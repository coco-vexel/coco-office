from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DEFAULT_OUTPUT = Path(__file__).resolve().parents[3] / "examples" / "output" / "zz_coco_native_word_line_grid_fixture.docx"


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Create a compact DOCX fixture for Word docGrid line-height validation.")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    build_document(doc)
    doc.save(args.out)
    patch_ooxml(args.out)
    print(args.out)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1.25)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.header_distance = Inches(0.39)
    section.footer_distance = Inches(0.39)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    set_east_asia_font(normal, "宋体")

    heading = doc.styles["Heading 1"]
    heading.font.name = "黑体"
    heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(8)
    set_east_asia_font(heading, "黑体")


def build_document(doc: Document) -> None:
    doc.add_heading("DocGrid 行距验证样本", level=1)
    add_marker_paragraph(doc, "CASE_AUTO_INHERIT")
    add_body_paragraphs(doc, "AUTO", count=16)

    doc.add_page_break()
    doc.add_heading("显式关闭 snapToGrid", level=1)
    add_marker_paragraph(doc, "CASE_SNAP_OFF")
    for paragraph in add_body_paragraphs(doc, "OFF", count=16):
        set_snap_to_grid(paragraph, False)

    doc.add_page_break()
    doc.add_heading("显式开启 snapToGrid", level=1)
    add_marker_paragraph(doc, "CASE_SNAP_ON")
    for paragraph in add_body_paragraphs(doc, "ON", count=16):
        set_snap_to_grid(paragraph, True)

    doc.add_page_break()
    doc.add_heading("表格内行距", level=1)
    add_marker_paragraph(doc, "CASE_TABLE")
    table = doc.add_table(rows=8, cols=3)
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            text = f"T{row_index + 1}-{cell_index + 1} 表格段落用于验证 docGrid 不应把单元格撑得过高。"
            paragraph = cell.paragraphs[0]
            paragraph.text = text
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)


def add_marker_paragraph(doc: Document, marker: str):
    paragraph = doc.add_paragraph(marker)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.size = Pt(9)
    return paragraph


def add_body_paragraphs(doc: Document, prefix: str, count: int):
    paragraphs = []
    base = (
        "生态流量泄放设施的方案验证需要稳定的正文行距、段落边界和分页行为。"
        "这一句包含中文、English words、数字 12345 以及标点，用于观察浏览器和 Word 的换行差异。"
    )
    for index in range(count):
        paragraph = doc.add_paragraph(f"{prefix}-{index + 1:02d} {base}")
        paragraph.paragraph_format.first_line_indent = Inches(0.33)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraphs.append(paragraph)
    return paragraphs


def set_east_asia_font(style_or_run, font_name: str) -> None:
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def set_snap_to_grid(paragraph, value: bool) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    snap = p_pr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        p_pr.append(snap)
    snap.set(qn("w:val"), "1" if value else "0")


def patch_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(dir=path.parent) as temp_dir:
        temp_path = Path(temp_dir) / path.name
        with ZipFile(path, "r") as source, ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/document.xml":
                    xml = data.decode("utf-8")
                    doc_grid = '<w:docGrid w:type="lines" w:linePitch="312" w:charSpace="0"/>'
                    if "<w:docGrid" in xml:
                        xml = re.sub(r"<w:docGrid\b[^>]*/>", doc_grid, xml)
                    else:
                        xml = xml.replace("</w:sectPr>", f"{doc_grid}</w:sectPr>")
                    data = xml.encode("utf-8")
                target.writestr(info, data)
        temp_path.replace(path)


if __name__ == "__main__":
    main()
